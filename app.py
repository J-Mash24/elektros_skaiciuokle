# ============================================================
# PV + BESS ENERGETINIO OPTIMIZAVIMO SKAIČIUOKLĖ
# Streamlit application
# ============================================================
import io
import numpy as np
import pandas as pd
import streamlit as st
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
import energy_model_core as core

# ============================================================
# 1. PROGRAMĖLĖS KONFIGŪRACIJA
# ============================================================

st.set_page_config(
    page_title="PV + BESS skaičiuoklė",
    layout="wide"
)

st.title("PV ir BESS energetinio optimizavimo skaičiuoklė")

st.caption(
    "Sezoninis PV generacijos, elektros vartojimo, BESS valdymo, "
    "kainų arbitražo ir 24 val. autonomijos modelis."
)

MODEL_CONFIG = core.ModelConfig(
    latitude=55.89,
    longitude=23.36,
    timezone="Europe/Vilnius",
    summer_months=(4, 5, 6, 7, 8, 9, 10),
    transition_months=(3, 11),
    winter_months=(12, 1, 2),
    low_price_quantile=0.25,
    high_price_quantile=0.75,
    bess_energy_capex_eur_kwh=300.0,
    bess_cycle_life=6000.0,
    degradation_cost_eur_kwh=None,
    minimum_arbitrage_margin_eur_kwh=0.01,
)


# ============================================================
# 2. PAGALBINĖS FUNKCIJOS
# ============================================================

def solar_shape(hour, sunrise=6, sunset=20):
    """
    Supaprastinta PV generacijos paros kreivė.
    """

    if hour < sunrise or hour >= sunset:
        return 0.0

    x = (
        (hour - sunrise)
        / (sunset - sunrise)
        * np.pi
    )

    return max(np.sin(x), 0.0)


def create_solar_weights():
    """
    Sukuria 24 valandų PV generacijos svorius.
    """

    hours = np.arange(24)

    weights = np.array(
        [
            solar_shape(hour)
            for hour in hours
        ],
        dtype=float
    )

    if weights.sum() == 0:
        raise ValueError(
            "Nepavyko suformuoti saulės generacijos profilio."
        )

    return weights / weights.sum()


def read_reference_pv_files(
    uploaded_files,
    reference_pv_kw
):
    """
    Perskaito vieną arba kelis referencinės PV elektrinės Excel failus.

    Programa automatiškai ieško Excel lapo, kuriame yra:
    - Statistical Period
    - PV Yield (kWh)

    Tikrinamos ir 1, ir 2 Excel antraštės eilutės.
    """

    frames = []

    required_columns = [
        "Statistical Period",
        "PV Yield (kWh)"
    ]

    for uploaded_file in uploaded_files:

        uploaded_file.seek(0)

        excel = pd.ExcelFile(uploaded_file)

        matching_df = None
        matching_sheet = None

        # Tikriname visus Excel lapus
        for sheet in excel.sheet_names:

            # Tikriname header=0 ir header=1
            for header_row in [0, 1]:

                uploaded_file.seek(0)

                try:

                    candidate = pd.read_excel(
                        uploaded_file,
                        sheet_name=sheet,
                        header=header_row
                    )

                except Exception:
                    continue

                if all(
                    col in candidate.columns
                    for col in required_columns
                ):

                    matching_df = candidate
                    matching_sheet = sheet
                    break

            if matching_df is not None:
                break

        if matching_df is None:
            # Failas nėra PV generacijos failas – praleidžiame
            continue

        temp = matching_df[
            [
                "Statistical Period",
                "PV Yield (kWh)"
            ]
        ].copy()

        temp.columns = [
            "date",
            "pv_ref_kwh"
        ]

        temp["source_file"] = uploaded_file.name
        temp["source_sheet"] = matching_sheet

        frames.append(temp)

    if len(frames) == 0:
        raise ValueError(
            "Tarp įkeltų failų nerasta nė vieno tinkamo "
            "PV generacijos failo. "
            "PV failuose turi būti stulpeliai "
            "'Statistical Period' ir 'PV Yield (kWh)'. "
            "ESO vartojimo ar sąskaitos failus kelkite į jiems skirtus laukus."
        )

    pv = pd.concat(
        frames,
        ignore_index=True
    )

    pv["date"] = pd.to_datetime(
        pv["date"],
        errors="coerce"
    )

    pv["pv_ref_kwh"] = pd.to_numeric(
        pv["pv_ref_kwh"],
        errors="coerce"
    )

    pv = pv.dropna(
        subset=[
            "date",
            "pv_ref_kwh"
        ]
    )

    pv = (
        pv.groupby(
            "date",
            as_index=False
        )["pv_ref_kwh"]
        .sum()
    )

    pv = pv.sort_values(
        "date"
    ).reset_index(drop=True)

    pv["yield_kwh_per_kw"] = (
        pv["pv_ref_kwh"]
        / reference_pv_kw
    )

    return pv


def create_hourly_pv_profile(pv_daily):
    """Dieninį referencinės PV profilį paverčia valandiniu pagal realią dienos trukmę."""
    return core.create_dynamic_hourly_pv_profile(
        pv_daily,
        cfg=MODEL_CONFIG
    )

def map_pv_profile_to_load_dates(hourly_pv, load_df):
    """PV profilį susieja su faktinio vartojimo datomis ir prideda sezono/dienos žymas."""
    return core.map_pv_to_load_dates(
        hourly_pv,
        load_df,
        cfg=MODEL_CONFIG
    )


def create_synthetic_load(
    datetime_series,
    annual_load_kwh
):
    """
    Sukuria sintetinį valandinį vartojimo profilį.
    """

    hour_factors = {
        0: 0.60, 1: 0.55, 2: 0.55, 3: 0.55,
        4: 0.55, 5: 0.60, 6: 0.70, 7: 0.85,
        8: 1.00, 9: 1.10, 10: 1.15, 11: 1.20,
        12: 1.20, 13: 1.15, 14: 1.10, 15: 1.05,
        16: 1.00, 17: 0.95, 18: 0.90, 19: 0.85,
        20: 0.80, 21: 0.75, 22: 0.70, 23: 0.65
    }

    month_factors = {
        1: 1.10,
        2: 1.08,
        3: 1.03,
        4: 0.98,
        5: 0.95,
        6: 0.93,
        7: 0.93,
        8: 0.95,
        9: 0.98,
        10: 1.02,
        11: 1.06,
        12: 1.10
    }

    df = pd.DataFrame({
        "datetime": pd.to_datetime(datetime_series)
    })

    df["hour"] = df["datetime"].dt.hour
    df["weekday"] = df["datetime"].dt.weekday
    df["month"] = df["datetime"].dt.month

    df["hour_factor"] = df["hour"].map(hour_factors)

    df["day_factor"] = np.where(
        df["weekday"] < 5,
        1.0,
        0.75
    )

    df["month_factor"] = (
        df["month"].map(month_factors)
    )

    df["raw_factor"] = (
        df["hour_factor"]
        * df["day_factor"]
        * df["month_factor"]
    )

    df["load_kwh"] = (
        df["raw_factor"]
        / df["raw_factor"].sum()
        * annual_load_kwh
    )

    return df[
        [
            "datetime",
            "load_kwh"
        ]
    ]

def read_actual_load(uploaded_file, sheet_name=None):
    """Perskaito aktyviosios P+ faktinės energijos vartojimo profilį."""
    uploaded_file.seek(0)

    if uploaded_file.name.lower().endswith(".csv"):
        df = pd.read_csv(uploaded_file)
    else:
        excel = pd.ExcelFile(uploaded_file)
        if sheet_name is None:
            matching_sheets = [
                sheet for sheet in excel.sheet_names
                if "valandin" in sheet.lower()
            ]
            sheet_name = matching_sheets[0] if matching_sheets else excel.sheet_names[0]
        uploaded_file.seek(0)
        df = pd.read_excel(uploaded_file, sheet_name=sheet_name)

    df.columns = [str(col).strip() for col in df.columns]

    # ESO faile kaupikliui aktuali aktyvioji vartojama energija P+.
    if "Energijos tipas" in df.columns:
        mask = df["Energijos tipas"].astype(str).str.strip().eq("P+")
        if mask.any():
            df = df[mask].copy()

    if "Suvartojimas" in df.columns:
        factual = (
            df["Suvartojimas"].astype(str).str.strip().str.lower().eq("faktinis")
        )
        if factual.any():
            df = df[factual].copy()

    datetime_candidates = ["datetime", "Data, valanda", "Data", "Laikotarpis"]
    load_candidates = ["load_kwh", "Kiekis, kWh", "Suvartojimas, kWh"]

    datetime_col = next((c for c in datetime_candidates if c in df.columns), None)
    load_col = next((c for c in load_candidates if c in df.columns), None)

    if datetime_col is None:
        raise ValueError(
            "Nepavyko rasti datos / laiko stulpelio. "
            f"Faile yra stulpeliai: {list(df.columns)}"
        )
    if load_col is None:
        raise ValueError(
            "Nepavyko rasti suvartojimo stulpelio. "
            f"Faile yra stulpeliai: {list(df.columns)}"
        )

    result = df[[datetime_col, load_col]].copy()
    result.columns = ["datetime", "load_kwh"]

    # Išsaugome vietinį ESO laiką ir nenaikiname valandos dėl UTC konversijos.
    datetime_text = result["datetime"].astype(str).str.strip()
    datetime_text = datetime_text.str.replace(
        r"([+-]\d{2}:?\d{2}|Z)$",
        "",
        regex=True
    )
    result["datetime"] = pd.to_datetime(datetime_text, errors="coerce")
    result["load_kwh"] = pd.to_numeric(result["load_kwh"], errors="coerce")
    result = result.dropna(subset=["datetime", "load_kwh"])
    result = (
        result.groupby("datetime", as_index=False)["load_kwh"]
        .sum()
        .sort_values("datetime")
        .reset_index(drop=True)
    )
    return result


# ============================================================
# ELEKTROS KAINŲ / SĄSKAITOS DUOMENŲ FUNKCIJOS
# ============================================================

def _clean_column_name(value):
    """Suvienodina Excel stulpelių pavadinimus."""
    return " ".join(str(value).replace("\n", " ").split()).strip()


def _find_header_row(uploaded_file, sheet_name, required_terms, max_rows=40):
    """
    Randa lentelės antraštės eilutę Excel lape.
    Naudinga sąskaitų failams, kuriuose lentelė prasideda ne 1 eilutėje.
    """
    uploaded_file.seek(0)
    raw = pd.read_excel(
        uploaded_file,
        sheet_name=sheet_name,
        header=None,
        nrows=max_rows
    )

    required_terms = [term.lower() for term in required_terms]

    for idx, row in raw.iterrows():
        cells = [_clean_column_name(v).lower() for v in row.tolist()]
        joined = " | ".join(cells)
        if all(term in joined for term in required_terms):
            return int(idx)

    return None


def read_market_price_file(uploaded_file, sheet_name=None):
    """
    Perskaito NPS / sąskaitos 15 min. arba kitokio intervalo duomenis.

    Tikimasi rasti bent:
    - Laikotarpis
    - Automatizuoti NPS LT, Eur/kWh

    Papildomai, jei yra, perskaitoma:
    - Suvartojimas, kWh
    - Priskaičiuota už el. energiją, Eur
    """
    uploaded_file.seek(0)
    excel = pd.ExcelFile(uploaded_file)

    candidate_sheets = [sheet_name] if sheet_name else excel.sheet_names
    selected_sheet = None
    header_row = None

    for sheet in candidate_sheets:
        if sheet is None:
            continue
        row = _find_header_row(
            uploaded_file,
            sheet,
            required_terms=["laikotarpis", "nps"]
        )
        if row is not None:
            selected_sheet = sheet
            header_row = row
            break

    if selected_sheet is None:
        raise ValueError(
            "Nepavyko rasti lapo su 'Laikotarpis' ir NPS kainos stulpeliu. "
            f"Rasti lapai: {excel.sheet_names}"
        )

    uploaded_file.seek(0)
    df = pd.read_excel(
        uploaded_file,
        sheet_name=selected_sheet,
        header=header_row
    )

    df.columns = [_clean_column_name(c) for c in df.columns]

    period_col = next(
        (c for c in df.columns if c.lower() == "laikotarpis"),
        None
    )
    price_col = next(
        (
            c for c in df.columns
            if "nps" in c.lower()
            and ("eur/kwh" in c.lower() or "eur / kwh" in c.lower())
        ),
        None
    )
    load_col = next(
        (c for c in df.columns if "suvartojimas" in c.lower() and "kwh" in c.lower()),
        None
    )
    charge_col = next(
        (c for c in df.columns if "priskai" in c.lower() and "energ" in c.lower() and "eur" in c.lower()),
        None
    )

    if period_col is None or price_col is None:
        raise ValueError(
            "Nerasti būtini kainų duomenų stulpeliai. "
            f"Rasti stulpeliai: {list(df.columns)}"
        )

    out = pd.DataFrame()

    period_text = df[period_col].astype(str).str.strip()
    start_text = period_text.str.extract(r"^\s*(.*?)\s+-\s+", expand=False)
    start_text = start_text.fillna(period_text)

    out["datetime"] = pd.to_datetime(
        start_text,
        errors="coerce"
    )

    out["buy_price_eur_kwh"] = pd.to_numeric(
        df[price_col],
        errors="coerce"
    )

    if load_col is not None:
        out["billed_load_kwh"] = pd.to_numeric(
            df[load_col],
            errors="coerce"
        )

    if charge_col is not None:
        out["energy_charge_eur"] = pd.to_numeric(
            df[charge_col],
            errors="coerce"
        )

    out = out.dropna(
        subset=["datetime", "buy_price_eur_kwh"]
    )

    # Išlaikome vietinį laiką, jei Excel reikšmė turi timezone.
    if out["datetime"].dt.tz is not None:
        out["datetime"] = out["datetime"].dt.tz_localize(None)

    out = out.sort_values("datetime").reset_index(drop=True)

    return out, selected_sheet


def align_price_to_base(base_df, price_df, fallback_price=0.15):
    """
    Priderina kainų intervalą prie energetinio modelio intervalo.
    Jei kainos trūksta, naudojama vartotojo nurodyta pakaitinė kaina.
    """
    result = base_df.copy()

    if price_df is None or len(price_df) == 0:
        result["buy_price_eur_kwh"] = float(fallback_price)
        return result, 0.0

    model_dt = detect_timestep_hours(result)
    price_dt = detect_timestep_hours(price_df)

    # Kainas agreguojame iki modelio intervalo, jei jos smulkesnės.
    if model_dt >= 1:
        freq = f"{int(round(model_dt * 60))}min"
    else:
        freq = f"{max(1, int(round(model_dt * 60)))}min"

    prices = price_df[["datetime", "buy_price_eur_kwh"]].copy()
    prices["period"] = prices["datetime"].dt.floor(freq)
    prices = (
        prices.groupby("period", as_index=False)["buy_price_eur_kwh"]
        .mean()
        .rename(columns={"period": "price_period"})
    )

    result["price_period"] = result["datetime"].dt.floor(freq)
    result = result.merge(
        prices,
        left_on="price_period",
        right_on="price_period",
        how="left"
    )

    coverage = float(result["buy_price_eur_kwh"].notna().mean() * 100)
    result["buy_price_eur_kwh"] = result["buy_price_eur_kwh"].fillna(float(fallback_price))
    result = result.drop(columns=["price_period"])

    return result, coverage

# ============================================================
# 3. BESS SIMULIAVIMO FUNKCIJA
# ============================================================
def detect_timestep_hours(df):
    """Automatiškai nustato modelio intervalo trukmę valandomis."""
    return core.detect_timestep_hours(df)

def simulate_bess(
    data,
    bess_kwh,
    bess_kw,
    charge_eff,
    discharge_eff,
    soc_min,
    soc_max,
    soc_initial,
    grid_import_limit_kw,
    grid_export_limit_kw
):
    """Sezoninis PV + BESS dispatch: vasara, pereinamasis laikotarpis ir žiema."""
    return core.simulate_seasonal_bess(
        data=data,
        bess_kwh=bess_kwh,
        bess_kw=bess_kw,
        charge_eff=charge_eff,
        discharge_eff=discharge_eff,
        soc_min=soc_min,
        soc_max=soc_max,
        soc_initial=soc_initial,
        grid_import_limit_kw=grid_import_limit_kw,
        grid_export_limit_kw=grid_export_limit_kw,
        cfg=MODEL_CONFIG
    )


# ============================================================
# 4. VIENO SCENARIJAUS FUNKCIJA
# ============================================================

def run_scenario(
    base_hourly,
    pv_kw,
    bess_kwh,
    bess_kw,
    charge_eff,
    discharge_eff,
    soc_min,
    soc_max,
    soc_initial,
    grid_import_limit_kw,
    grid_export_limit_kw
):
    """Paleidžia naują sezoninį modelio branduolį."""
    return core.run_seasonal_scenario(
        base_data=base_hourly,
        pv_kw=pv_kw,
        bess_kwh=bess_kwh,
        bess_kw=bess_kw,
        charge_eff=charge_eff,
        discharge_eff=discharge_eff,
        soc_min=soc_min,
        soc_max=soc_max,
        soc_initial=soc_initial,
        grid_import_limit_kw=grid_import_limit_kw,
        grid_export_limit_kw=grid_export_limit_kw,
        cfg=MODEL_CONFIG
    )


# ============================================================
# 5. EKONOMINĖ FUNKCIJA
# ============================================================

def add_economics(
    scenarios,
    pv_cost_eur_kw,
    bess_cost_eur_kwh,
    bess_cost_eur_kw,
    electricity_buy_price,
    electricity_sell_price,
    pv_opex_rate,
    bess_opex_rate,
    discount_rate,
    project_years,
    annual_load_kwh
):

    df = scenarios.copy()

    df["pv_capex_eur"] = (
        df["pv_kw"]
        * pv_cost_eur_kw
    )

    df["bess_capex_eur"] = (
        df["bess_kwh"]
        * bess_cost_eur_kwh
        +
        df["bess_kw"]
        * bess_cost_eur_kw
    )

    df["total_capex_eur"] = (
        df["pv_capex_eur"]
        +
        df["bess_capex_eur"]
    )

    if (
        "actual_grid_cost_eur" in df.columns
        and df["actual_grid_cost_eur"].notna().any()
    ):
        df["grid_cost_eur"] = df["actual_grid_cost_eur"].fillna(
            df["grid_import_kwh"] * electricity_buy_price
        )
    else:
        df["grid_cost_eur"] = (
            df["grid_import_kwh"]
            * electricity_buy_price
        )

    df["export_revenue_eur"] = (
        df["export_kwh"]
        * electricity_sell_price
    )

    df["opex_eur"] = (
        df["pv_capex_eur"]
        * pv_opex_rate
        +
        df["bess_capex_eur"]
        * bess_opex_rate
    )

    df["annual_cost_eur"] = (
        df["grid_cost_eur"]
        +
        df["opex_eur"]
        -
        df["export_revenue_eur"]
    )

    if (
        "baseline_energy_cost_eur" in df.columns
        and df["baseline_energy_cost_eur"].notna().any()
    ):
        base_annual_cost = float(
            df["baseline_energy_cost_eur"].dropna().iloc[0]
        )
    else:
        base_annual_cost = (
            annual_load_kwh
            * electricity_buy_price
        )

    df["annual_savings_eur"] = (
        base_annual_cost
        -
        df["annual_cost_eur"]
    )

    df["simple_payback_years"] = np.where(
        df["annual_savings_eur"] > 0,
        (
            df["total_capex_eur"]
            / df["annual_savings_eur"]
        ),
        np.nan
    )

    discount_factor = sum(
        1
        / (
            1 + discount_rate
        ) ** year
        for year in range(
            1,
            project_years + 1
        )
    )

    df["npc_eur"] = (
        df["total_capex_eur"]
        +
        df["annual_cost_eur"]
        * discount_factor
    )

    df["npv_eur"] = (
        -df["total_capex_eur"]
        +
        df["annual_savings_eur"]
        * discount_factor
    )

    return df


# ============================================================
# 6. OPTIMIZAVIMO FUNKCIJA
# ============================================================

def optimize_system(
    base_hourly,
    pv_sizes,
    bess_energy_sizes,
    bess_power_sizes,
    charge_eff,
    discharge_eff,
    soc_min,
    soc_max,
    soc_initial,
    grid_import_limit_kw,
    grid_export_limit_kw
):

    scenario_results = []

    total = (
        len(pv_sizes)
        * len(bess_energy_sizes)
        * len(bess_power_sizes)
    )

    progress = st.progress(0)

    counter = 0

    for pv_kw in pv_sizes:

        for bess_kwh in bess_energy_sizes:

            for bess_kw in bess_power_sizes:

                counter += 1

                progress.progress(
                    min(
                        counter / total,
                        1.0
                    )
                )

                # BESS = 0 / 0 leidžiamas
                if (
                    bess_kwh == 0
                    and bess_kw != 0
                ):
                    continue

                if (
                    bess_kwh > 0
                    and bess_kw == 0
                ):
                    continue

                summary, _ = run_scenario(
                    base_hourly=base_hourly,
                    pv_kw=pv_kw,
                    bess_kwh=bess_kwh,
                    bess_kw=bess_kw,
                    charge_eff=charge_eff,
                    discharge_eff=discharge_eff,
                    soc_min=soc_min,
                    soc_max=soc_max,
                    soc_initial=soc_initial,
                    grid_import_limit_kw=grid_import_limit_kw,
                    grid_export_limit_kw=grid_export_limit_kw
                )

                scenario_results.append(
                    summary
                )

    progress.empty()

    return pd.DataFrame(
        scenario_results
    )


# ============================================================
# 7. ŠONINĖ JUOSTA - PAGRINDINIAI PARAMETRAI
# ============================================================

st.sidebar.header(
    "Objekto parametrai"
)

annual_load_kwh = st.sidebar.number_input(
    "Metinis elektros suvartojimas, kWh",
    min_value=0.0,
    value=659955.0,
    step=1000.0
)

model_year = st.sidebar.number_input(
    "Modeliuojamo profilio metai",
    min_value=2000,
    max_value=2100,
    value=2025,
    step=1,
    help="Naudojama tik tada, kai vartojimas modeliuojamas, o ne įkeliamas iš failo."
)

reference_pv_kw = st.sidebar.number_input(
    "Referencinės PV elektrinės galia, kW",
    min_value=0.01,
    value=21.12,
    step=0.01,
    help="Aktualu tik pasirinkus referencinių PV failų režimą."
)

grid_import_limit_kw = st.sidebar.number_input(
    "Importo galios riba, kW",
    min_value=0.0,
    value=750.0,
    step=10.0
)

grid_export_limit_kw = st.sidebar.number_input(
    "Eksporto galios riba, kW",
    min_value=0.0,
    value=1000.0,
    step=10.0
)


# ============================================================
# 8. PAGRINDINIAI SKIRTUKAI
# ============================================================

tab_data, tab_scenario, tab_optimization, tab_results = st.tabs(
    [
        "1. Duomenys",
        "2. Scenarijus",
        "3. Optimizavimas",
        "4. Rezultatai"
    ]
)



def _flatten_dict(prefix, value, rows):
    """Paverčia įdėtą dict struktūrą į Excel draugišką rodiklių lentelę."""
    if isinstance(value, dict):
        for key, subvalue in value.items():
            child = f"{prefix}.{key}" if prefix else str(key)
            _flatten_dict(child, subvalue, rows)
    elif isinstance(value, pd.DataFrame):
        return
    else:
        rows.append({"Rodiklis": prefix, "Reikšmė": value})


def _autosize_excel_sheet(ws, max_width=45):
    for col_idx, column_cells in enumerate(ws.columns, start=1):
        max_len = 0
        for cell in column_cells:
            if cell.value is None:
                continue
            max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max(max_len + 2, 10), max_width)


def _format_excel_workbook(writer):
    """Lengvas, profesionalus Excel formatavimas eksportuojamoms lentelėms."""
    wb = writer.book
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    section_fill = PatternFill("solid", fgColor="D9EAF7")

    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        ws.sheet_view.showGridLines = False

        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        _autosize_excel_sheet(ws)

        # Bendriniai skaičių formatai pagal stulpelių pavadinimus.
        headers = {cell.column: str(cell.value or "") for cell in ws[1]}
        for col_idx, header in headers.items():
            h = header.lower()
            for row in range(2, ws.max_row + 1):
                cell = ws.cell(row=row, column=col_idx)
                if isinstance(cell.value, (int, float)):
                    if "pct" in h or "%" in h or "procent" in h:
                        cell.number_format = '0.0'
                    elif "eur" in h or "€" in h or "cost" in h or "capex" in h or "npc" in h or "npv" in h:
                        cell.number_format = '#,##0.00;[Red](#,##0.00);-'
                    elif "kwh" in h or "kw" in h or "mwh" in h:
                        cell.number_format = '#,##0.00;[Red](#,##0.00);-'
                    else:
                        cell.number_format = '#,##0.00;[Red](#,##0.00);-'


def build_excel_export(
    summary,
    hourly_result,
    monthly,
    result_sizing,
    daylight,
    assumptions,
    optimization_results=None,
    best_scenario=None,
    validation=None,
):
    """Sukuria vieną .xlsx failą su visais modelio rezultatais."""
    output = io.BytesIO()

    summary_rows = []
    for key, value in summary.items():
        summary_rows.append({"Rodiklis": key, "Reikšmė": value})
    summary_df = pd.DataFrame(summary_rows)

    sizing_rows = []
    for key, value in result_sizing.items():
        if key == "monthly_daylight":
            continue
        _flatten_dict(key, value, sizing_rows)
    sizing_df = pd.DataFrame(sizing_rows)

    assumptions_df = pd.DataFrame(
        [{"Parametras": key, "Reikšmė": value} for key, value in assumptions.items()]
    )

    monthly_export = monthly.reset_index().copy()
    daylight_export = daylight.reset_index().copy()

    # Sezoninė suvestinė tiesiai iš valandinių rezultatų.
    seasonal_columns = [
        c for c in [
            "load_kwh", "pv_kwh", "grid_import_kwh", "grid_import_load_kwh",
            "grid_charge_kwh", "export_kwh", "bess_charge_kwh",
            "bess_discharge_kwh", "curtailment_kwh", "unserved_load_kwh"
        ] if c in hourly_result.columns
    ]
    if "season" in hourly_result.columns and seasonal_columns:
        seasonal_df = (
            hourly_result.groupby("season", as_index=False)[seasonal_columns]
            .sum(numeric_only=True)
        )
    else:
        seasonal_df = pd.DataFrame()

    validation_df = pd.DataFrame(
        [{"Patikra": k, "Reikšmė": v} for k, v in (validation or {}).items()]
    )

    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        summary_df.to_excel(writer, sheet_name="Summary", index=False)
        assumptions_df.to_excel(writer, sheet_name="Assumptions", index=False)
        sizing_df.to_excel(writer, sheet_name="BESS sizing", index=False)
        daylight_export.to_excel(writer, sheet_name="Daylight", index=False)
        monthly_export.to_excel(writer, sheet_name="Monthly", index=False)
        if not seasonal_df.empty:
            seasonal_df.to_excel(writer, sheet_name="Seasonal", index=False)
        hourly_result.to_excel(writer, sheet_name="Hourly results", index=False)
        if not validation_df.empty:
            validation_df.to_excel(writer, sheet_name="Validation", index=False)

        if optimization_results is not None and len(optimization_results) > 0:
            optimization_results.to_excel(writer, sheet_name="Optimization", index=False)
            top20_export = optimization_results.sort_values("npc_eur").head(20)
            top20_export.to_excel(writer, sheet_name="Top scenarios", index=False)

        if best_scenario is not None:
            best_df = pd.DataFrame(
                [{"Rodiklis": k, "Reikšmė": v} for k, v in dict(best_scenario).items()]
            )
            best_df.to_excel(writer, sheet_name="Best scenario", index=False)

        _format_excel_workbook(writer)

    output.seek(0)
    return output.getvalue()


# ============================================================
# 9. DUOMENŲ SKIRTUKAS
# ============================================================

with tab_data:

    st.subheader("PV generacijos duomenys")

    pv_source = st.radio(
        "PV generacijos šaltinis",
        [
            "Modeliuojamas PV profilis (nereikia failo)",
            "Referencinės PV elektrinės failai"
        ],
        horizontal=True
    )

    annual_specific_yield = st.number_input(
        "Metinė specifinė PV generacija, kWh/kW",
        min_value=100.0,
        max_value=2000.0,
        value=904.5,
        step=10.0,
        help=(
            "Naudojama modeliuojamam PV profiliui. Numatytoji 904,5 kWh/kW reikšmė "
            "kalibruota pagal šiame projekte naudotos 21,12 kW referencinės elektrinės 2025 m. duomenis."
        )
    )

    pv_files = []

    if pv_source == "Referencinės PV elektrinės failai":
        pv_files = st.file_uploader(
            "Įkelkite referencinės PV elektrinės Excel failus",
            type=["xlsx"],
            accept_multiple_files=True
        )

        st.info(
            "Šis režimas naudoja faktinį referencinės elektrinės profilį. "
            "Jei failų neįkelsite, modelis nebus blokuojamas – galite pasirinkti "
            "„Modeliuojamas PV profilis (nereikia failo)“."
        )

        if pv_files:
            view_mode = st.radio(
                "Excel failų peržiūra",
                ["Failai atskirai", "Sujungti PV duomenys"],
                horizontal=True
            )

            if view_mode == "Failai atskirai":
                st.write(f"Įkelta failų: {len(pv_files)}")
                for file_index, uploaded_file in enumerate(pv_files):
                    with st.expander(
                        f"{file_index + 1}. {uploaded_file.name}",
                        expanded=False
                    ):
                        try:
                            uploaded_file.seek(0)
                            excel_file = pd.ExcelFile(uploaded_file)
                            selected_sheet = st.selectbox(
                                "Pasirinkite Excel lapą",
                                excel_file.sheet_names,
                                key=f"pv_sheet_{file_index}_{uploaded_file.name}"
                            )
                            uploaded_file.seek(0)
                            full_excel = pd.read_excel(
                                uploaded_file,
                                sheet_name=selected_sheet
                            )
                            st.dataframe(
                                full_excel,
                                use_container_width=True,
                                height=500
                            )
                        except Exception as exc:
                            st.error(
                                f"Nepavyko perskaityti failo {uploaded_file.name}: {exc}"
                            )
            else:
                try:
                    for uploaded_file in pv_files:
                        uploaded_file.seek(0)
                    pv_daily_preview = read_reference_pv_files(
                        pv_files,
                        reference_pv_kw
                    )
                    reference_generation = pv_daily_preview["pv_ref_kwh"].sum()
                    specific_yield = reference_generation / reference_pv_kw
                    c1, c2, c3 = st.columns(3)
                    c1.metric("Referencinė PV galia", f"{reference_pv_kw:.2f} kW")
                    c2.metric(
                        "Referencinė generacija",
                        f"{reference_generation / 1000:.2f} MWh"
                    )
                    c3.metric(
                        "Specifinė generacija",
                        f"{specific_yield:.1f} kWh/kW"
                    )
                    st.dataframe(
                        pv_daily_preview,
                        use_container_width=True,
                        height=500
                    )
                except Exception as exc:
                    st.warning(f"PV duomenų patikros pranešimas: {exc}")
    else:
        st.success(
            "PV failas nėra būtinas. Modelis sugeneruos sezoninį PV profilį pagal "
            "metinę specifinę generaciją, mėnesinį pasiskirstymą ir realią dienos trukmę."
        )

    # ========================================================
    # ELEKTROS VARTOJIMO DUOMENYS
    # ========================================================

    st.subheader("Elektros vartojimo duomenys")

    load_source = st.radio(
        "Vartojimo profilis",
        ["Modeliuotas profilis", "Faktinis profilis"]
    )

    load_file = None
    load_sheet = None
    load_preview = None

    if load_source == "Faktinis profilis":
        load_file = st.file_uploader(
            "Įkelkite faktinio elektros vartojimo CSV arba Excel failą",
            type=["csv", "xlsx"],
            key="actual_load_file"
        )

        if load_file is not None:
            if load_file.name.lower().endswith(".xlsx"):
                load_file.seek(0)
                excel = pd.ExcelFile(load_file)
                load_sheet = st.selectbox(
                    "Pasirinkite vartojimo duomenų Excel lapą",
                    excel.sheet_names,
                    key="load_sheet_selector"
                )
                st.caption("Aptikti lapai: " + ", ".join(excel.sheet_names))
                try:
                    load_file.seek(0)
                    load_preview = pd.read_excel(
                        load_file,
                        sheet_name=load_sheet
                    )
                    st.dataframe(
                        load_preview,
                        use_container_width=True,
                        height=450
                    )
                except Exception as exc:
                    st.error(f"Nepavyko parodyti Excel lapo: {exc}")
            else:
                load_file.seek(0)
                load_preview = pd.read_csv(load_file)
                st.dataframe(
                    load_preview,
                    use_container_width=True,
                    height=450
                )

    # ========================================================
    # ELEKTROS KAINŲ / SĄSKAITOS DUOMENYS
    # ========================================================

    st.subheader("Elektros kainų / sąskaitos duomenys")

    price_source = st.radio(
        "Elektros pirkimo kainos šaltinis",
        ["Fiksuota kaina", "Faktinės NPS kainos"],
        horizontal=True
    )

    price_file = None
    price_sheet = None
    price_preview = None

    price_fallback = st.number_input(
        "Pakaitinė pirkimo kaina, kai faktinės kainos nėra, €/kWh",
        min_value=0.0,
        value=0.15,
        format="%.4f"
    )

    if price_source == "Faktinės NPS kainos":
        price_file = st.file_uploader(
            "Įkelkite NPS / elektros sąskaitos Excel failą",
            type=["xlsx"],
            key="market_price_file"
        )
        st.info(
            "Programa ieškos lentelės su stulpeliais 'Laikotarpis' ir "
            "'Automatizuoti NPS LT, Eur/kWh'."
        )
        if price_file is not None:
            price_file.seek(0)
            price_excel = pd.ExcelFile(price_file)
            price_sheet = st.selectbox(
                "Pasirinkite kainų / sąskaitos Excel lapą",
                price_excel.sheet_names,
                key="price_sheet_selector"
            )
            try:
                parsed_price_preview, detected_price_sheet = read_market_price_file(
                    price_file,
                    sheet_name=price_sheet
                )
                st.caption(
                    f"Atpažintas lapas: {detected_price_sheet} | "
                    f"Įrašų: {len(parsed_price_preview)}"
                )
                st.dataframe(
                    parsed_price_preview,
                    use_container_width=True,
                    height=400
                )
            except Exception as exc:
                st.error(
                    f"Nepavyko perskaityti kainų / sąskaitos duomenų: {exc}"
                )

# ============================================================
# 10. BAZINIO PROFILIO PARUOŠIMAS
# ============================================================

base_hourly = None
load_df = None
data_ready = False

actual_load_total = None
first_timestamp = None
last_timestamp = None
dt_hours = None
peak_load_kw = None
price_df = None
price_coverage_pct = None
data_error_message = None
pv_profile_source_used = None

try:
    # --------------------------------------------------------
    # 1. VARTOJIMO PROFILIS
    # --------------------------------------------------------
    if load_source == "Faktinis profilis":
        if load_file is not None:
            load_df = read_actual_load(
                load_file,
                sheet_name=load_sheet
            )
    else:
        synthetic_index = pd.date_range(
            start=f"{int(model_year)}-01-01 00:00:00",
            end=f"{int(model_year)}-12-31 23:00:00",
            freq="h"
        )
        load_df = create_synthetic_load(
            synthetic_index,
            annual_load_kwh
        )

    # --------------------------------------------------------
    # 2. PV PROFILIS – FAILAS NEPRIVALOMAS
    # --------------------------------------------------------
    if load_df is not None and len(load_df) > 0:
        if pv_source == "Referencinės PV elektrinės failai" and pv_files:
            pv_daily = read_reference_pv_files(
                pv_files,
                reference_pv_kw
            )
            hourly_pv = create_hourly_pv_profile(pv_daily)
            base_hourly = map_pv_profile_to_load_dates(
                hourly_pv,
                load_df
            )
            pv_profile_source_used = "Referencinės PV elektrinės duomenys"
        else:
            # Jei pasirinktas modeliuojamas režimas arba referencinių failų nėra,
            # modelis automatiškai sukuria sezoninį PV profilį.
            base_hourly = core.create_synthetic_pv_for_load(
                load_df=load_df,
                annual_specific_yield_kwh_kw=float(annual_specific_yield),
                cfg=MODEL_CONFIG
            )
            pv_profile_source_used = "Modeliuojamas sezoninis PV profilis"

    # --------------------------------------------------------
    # 3. ELEKTROS KAINOS
    # --------------------------------------------------------
    if base_hourly is not None and len(base_hourly) > 0:
        if (
            price_source == "Faktinės NPS kainos"
            and price_file is not None
        ):
            price_df, _ = read_market_price_file(
                price_file,
                sheet_name=price_sheet
            )
            base_hourly, price_coverage_pct = align_price_to_base(
                base_hourly,
                price_df,
                fallback_price=price_fallback
            )
        else:
            base_hourly["buy_price_eur_kwh"] = float(price_fallback)
            price_coverage_pct = (
                100.0 if price_source == "Fiksuota kaina" else 0.0
            )

    # --------------------------------------------------------
    # 4. DUOMENŲ PATIKRA
    # --------------------------------------------------------
    if (
        load_df is not None
        and base_hourly is not None
        and len(base_hourly) > 0
    ):
        data_ready = True
        actual_load_total = load_df["load_kwh"].sum()
        first_timestamp = load_df["datetime"].min()
        last_timestamp = load_df["datetime"].max()
        dt_hours = detect_timestep_hours(load_df)
        peak_load_kw = load_df["load_kwh"].max() / dt_hours

except ValueError as exc:
    data_error_message = str(exc)
except Exception as exc:
    data_error_message = f"Duomenų paruošimo klaida: {exc}"

# ============================================================
# 11. SCENARIJAUS SKIRTUKAS
# ============================================================

with tab_scenario:

    st.subheader(
        "Vieno PV + BESS scenarijaus analizė"
    )

    if not data_ready:
        if data_error_message:
            st.warning(
                "Modelis dar neparuoštas. " + data_error_message
            )
        elif load_source == "Faktinis profilis" and load_file is None:
            st.warning(
                "Modelis dar neparuoštas: pasirinktas faktinis vartojimas, "
                "bet vartojimo failas neįkeltas."
            )
        else:
            st.warning("Modelis dar neparuoštas. Patikrinkite įkeltus duomenis.")

    else:

        col1, col2, col3 = st.columns(3)

        with col1:

            pv_kw = st.number_input(
                "PV elektrinės galia, kW",
                min_value=0.0,
                value=730.0,
                step=10.0
            )

        with col2:

            bess_kwh = st.number_input(
                "BESS talpa, kWh",
                min_value=0.0,
                value=1000.0,
                step=50.0
            )

        with col3:

            bess_kw = st.number_input(
                "BESS galia, kW",
                min_value=0.0,
                value=500.0,
                step=25.0
            )

        if bess_kw > 0:

            st.metric(
                "BESS nominali trukmė",
                f"{bess_kwh / bess_kw:.2f} h"
            )

        sizing = core.calculate_sizing_metrics(
            base_hourly,
            pv_kw=pv_kw,
            discharge_eff=0.95,
            soc_min=0.05,
            soc_max=0.95,
            summer_reserve_pct=15.0,
            cfg=MODEL_CONFIG
        )

        st.markdown("### Modelio techniniai orientyrai")
        s1, s2, s3, s4 = st.columns(4)
        s1.metric(
            "Vasaros BESS orientyras",
            f"{sizing['summer_bess_nominal_kwh'] / 1000:.2f} MWh"
        )
        s2.metric(
            "Žiemos techninis orientyras",
            f"{sizing['winter_bess_technical_kwh'] / 1000:.2f} MWh"
        )
        s3.metric(
            "24 h autonomija (P90)",
            f"{sizing['autonomy_bess_p90_kwh'] / 1000:.2f} MWh"
        )
        s4.metric(
            "24 h autonomija (maks.)",
            f"{sizing['autonomy_bess_max_kwh'] / 1000:.2f} MWh"
        )

        st.caption(
            "Vasara: balandis–spalis – PV dengia apkrovą, PV perteklius krauna BESS, "
            "BESS naudojamas naktį. Kovas ir lapkritis – mišrus režimas. "
            "Žiema: gruodis–vasaris – PV dengia dalį apkrovos, o BESS gali būti "
            "kraunamas iš tinklo tik tada, kai faktinės kainos rodo ekonomiškai "
            "pagrįstą arbitražo galimybę."
        )

        with st.expander(
            "Išplėstiniai BESS parametrai"
        ):

            charge_eff_pct = st.slider(
                "Įkrovimo efektyvumas, %",
                50.0,
                100.0,
                95.0
            )

            discharge_eff_pct = st.slider(
                "Iškrovimo efektyvumas, %",
                50.0,
                100.0,
                95.0
            )

            soc_min_pct = st.slider(
                "Minimalus SOC, %",
                0.0,
                50.0,
                5.0
            )

            soc_max_pct = st.slider(
                "Maksimalus SOC, %",
                50.0,
                100.0,
                95.0
            )

            soc_initial_pct = st.slider(
                "Pradinis SOC, %",
                0.0,
                100.0,
                5.0
            )

        if st.button(
            "Skaičiuoti scenarijų",
            type="primary"
        ):

            if soc_min_pct >= soc_max_pct:

                st.error(
                    "Minimalus SOC turi būti mažesnis už maksimalų SOC."
                )

            else:

                scenario_summary, scenario_hourly = run_scenario(
                    base_hourly=base_hourly,
                    pv_kw=pv_kw,
                    bess_kwh=bess_kwh,
                    bess_kw=bess_kw,
                    charge_eff=charge_eff_pct / 100,
                    discharge_eff=discharge_eff_pct / 100,
                    soc_min=soc_min_pct / 100,
                    soc_max=soc_max_pct / 100,
                    soc_initial=soc_initial_pct / 100,
                    grid_import_limit_kw=grid_import_limit_kw,
                    grid_export_limit_kw=grid_export_limit_kw
                )

                st.session_state[
                    "scenario_summary"
                ] = scenario_summary

                st.session_state[
                    "scenario_hourly"
                ] = scenario_hourly


# ============================================================
# 12. OPTIMIZAVIMO SKIRTUKAS
# ============================================================

with tab_optimization:

    st.subheader(
        "Automatinė PV + BESS optimizacija"
    )

    if not data_ready:
        st.warning(
            "Optimizavimas bus aktyvus, kai bus paruošti PV ir vartojimo duomenys."
        )

    else:

        st.markdown(
            "### Paieškos ribos"
        )

        c1, c2, c3 = st.columns(3)

        with c1:

            pv_min = st.number_input(
                "PV minimumas, kW",
                value=300,
                step=50
            )

            pv_max = st.number_input(
                "PV maksimumas, kW",
                value=1000,
                step=50
            )

            pv_step = st.number_input(
                "PV žingsnis, kW",
                value=50,
                min_value=1
            )

        with c2:

            bess_e_min = st.number_input(
                "BESS talpos minimumas, kWh",
                value=0,
                step=250
            )

            bess_e_max = st.number_input(
                "BESS talpos maksimumas, kWh",
                value=2000,
                step=250
            )

            bess_e_step = st.number_input(
                "BESS talpos žingsnis, kWh",
                value=250,
                min_value=1
            )

        with c3:

            bess_p_min = st.number_input(
                "BESS galios minimumas, kW",
                value=0,
                step=125
            )

            bess_p_max = st.number_input(
                "BESS galios maksimumas, kW",
                value=750,
                step=125
            )

            bess_p_step = st.number_input(
                "BESS galios žingsnis, kW",
                value=125,
                min_value=1
            )

        target_ssr = st.slider(
            "Minimalus energetinis savarankiškumas, %",
            min_value=0,
            max_value=100,
            value=60
        )

        autonomy_constraint = st.selectbox(
            "24 h autonomijos reikalavimas",
            [
                "Netaikyti",
                "P90 24 h poreikis",
                "Maksimalus 24 h poreikis"
            ],
            help=(
                "Konservatyvus autonomijos kriterijus skaičiuojamas be PV indėlio. "
                "Jei taikomas, optimizacijoje atmetami mažesnės talpos BESS variantai."
            )
        )

        autonomy_sizing = core.calculate_sizing_metrics(
            base_hourly,
            pv_kw=0.0,
            discharge_eff=0.95,
            soc_min=0.05,
            soc_max=0.95,
            summer_reserve_pct=15.0,
            cfg=MODEL_CONFIG
        )

        if autonomy_constraint == "P90 24 h poreikis":
            autonomy_min_bess = autonomy_sizing["autonomy_bess_p90_kwh"]
        elif autonomy_constraint == "Maksimalus 24 h poreikis":
            autonomy_min_bess = autonomy_sizing["autonomy_bess_max_kwh"]
        else:
            autonomy_min_bess = 0.0

        if autonomy_min_bess > 0:
            st.caption(
                f"Minimalus BESS pagal pasirinktą 24 h autonomijos kriterijų: "
                f"{autonomy_min_bess / 1000:.2f} MWh."
            )

        st.markdown(
            "### Ekonominės prielaidos"
        )

        e1, e2, e3 = st.columns(3)

        with e1:

            pv_cost = st.number_input(
                "PV CAPEX, €/kW",
                value=650.0
            )

            bess_energy_cost = st.number_input(
                "BESS CAPEX, €/kWh",
                value=300.0
            )

            bess_power_cost = st.number_input(
                "BESS galios CAPEX, €/kW",
                value=100.0
            )

        with e2:

            buy_price = st.number_input(
                "Elektros pirkimo kaina, €/kWh",
                value=float(price_fallback),
                format="%.4f",
                help=(
                    "Naudojama kaip fiksuota kaina arba kaip pakaitinė kaina "
                    "intervalams, kuriems nėra faktinės NPS kainos."
                )
            )

            sell_price = st.number_input(
                "Eksporto vertė, €/kWh",
                value=0.05,
                format="%.3f"
            )

        with e3:

            discount_rate_pct = st.number_input(
                "Diskonto norma, %",
                value=5.0
            )

            project_years = st.number_input(
                "Projekto laikotarpis, metai",
                value=20,
                min_value=1
            )

            pv_opex_pct = st.number_input(
                "PV OPEX, % CAPEX/metus",
                value=1.0
            )

            bess_opex_pct = st.number_input(
                "BESS OPEX, % CAPEX/metus",
                value=1.0
            )

        if st.button(
            "Paleisti optimizaciją",
            type="primary"
        ):

            pv_sizes = np.arange(
                pv_min,
                pv_max + pv_step,
                pv_step
            )

            bess_energy_sizes = np.arange(
                bess_e_min,
                bess_e_max + bess_e_step,
                bess_e_step
            )

            bess_power_sizes = np.arange(
                bess_p_min,
                bess_p_max + bess_p_step,
                bess_p_step
            )

            raw_scenarios = optimize_system(
                base_hourly=base_hourly,
                pv_sizes=pv_sizes,
                bess_energy_sizes=bess_energy_sizes,
                bess_power_sizes=bess_power_sizes,
                charge_eff=0.95,
                discharge_eff=0.95,
                soc_min=0.05,
                soc_max=0.95,
                soc_initial=0.05,
                grid_import_limit_kw=grid_import_limit_kw,
                grid_export_limit_kw=grid_export_limit_kw
            )

            # Atmetame techniškai netinkamus scenarijus
            valid = raw_scenarios[
                raw_scenarios[
                    "unserved_load_kwh"
                ] < 0.001
            ].copy()

            if autonomy_min_bess > 0:
                valid = valid[
                    valid["bess_kwh"] >= autonomy_min_bess
                ].copy()

                if len(valid) == 0:
                    st.warning(
                        "Pasirinktame optimizavimo diapazone nėra BESS talpos, "
                        "atitinkančios 24 h autonomijos kriterijų. Padidinkite BESS "
                        "talpos maksimumą."
                    )

            economic = add_economics(
                scenarios=valid,
                pv_cost_eur_kw=pv_cost,
                bess_cost_eur_kwh=bess_energy_cost,
                bess_cost_eur_kw=bess_power_cost,
                electricity_buy_price=buy_price,
                electricity_sell_price=sell_price,
                pv_opex_rate=pv_opex_pct / 100,
                bess_opex_rate=bess_opex_pct / 100,
                discount_rate=discount_rate_pct / 100,
                project_years=int(project_years),
                annual_load_kwh=base_hourly["load_kwh"].sum()
            )

            st.session_state[
                "optimization_results"
            ] = economic

            acceptable = economic[
                economic[
                    "self_sufficiency_pct"
                ] >= target_ssr
            ]

            if len(acceptable) > 0:

                best = acceptable.loc[
                    acceptable[
                        "npc_eur"
                    ].idxmin()
                ]

                st.session_state[
                    "best_scenario"
                ] = best

            else:

                st.session_state[
                    "best_scenario"
                ] = None


# ============================================================
# 13. REZULTATŲ SKIRTUKAS
# ============================================================

with tab_results:

    st.subheader(
        "Rezultatai"
    )

    # ========================================================
    # VIENO SCENARIJAUS REZULTATAI
    # ========================================================

    if "scenario_summary" in st.session_state:

        summary = st.session_state[
            "scenario_summary"
        ]

        hourly_result = st.session_state[
            "scenario_hourly"
        ]

        st.markdown(
            "## Pasirinktas scenarijus"
        )

        c1, c2, c3, c4 = st.columns(4)

        c1.metric(
            "PV generacija",
            f"{summary['pv_generation_kwh'] / 1000:.1f} MWh"
        )

        c2.metric(
            "Importas",
            f"{summary['grid_import_kwh'] / 1000:.1f} MWh"
        )

        c3.metric(
            "Savarankiškumas",
            f"{summary['self_sufficiency_pct']:.1f} %"
        )

        c4.metric(
            "PV panaudojimas",
            f"{summary['self_consumption_pct']:.1f} %"
        )

        c5, c6, c7, c8 = st.columns(4)

        c5.metric(
            "Eksportas",
            f"{summary['export_kwh'] / 1000:.1f} MWh"
        )

        c6.metric(
            "BESS iškrovimas",
            f"{summary['bess_discharge_kwh'] / 1000:.1f} MWh"
        )

        c7.metric(
            "BESS ciklai",
            f"{summary['equivalent_cycles']:.0f}"
        )

        c8.metric(
            "PV apribojimas",
            f"{summary['curtailment_kwh'] / 1000:.1f} MWh"
        )

        if summary.get("grid_charge_kwh", 0.0) > 0:
            st.info(
                f"Žiemos / pereinamuoju laikotarpiu iš tinklo į BESS įkrauta "
                f"{summary['grid_charge_kwh'] / 1000:.1f} MWh. "
                "Įkrovimas leidžiamas tik ekonomiškai pagrįstais pigios kainos intervalais."
            )

        # ----------------------------------------------------
        # VARTOJIMO DUOMENŲ PATIKRA
        # ----------------------------------------------------

        if (
            load_df is not None
            and actual_load_total is not None
        ):

            st.markdown(
                "### Vartojimo duomenų patikra"
            )

            c1, c2, c3, c4 = st.columns(4)

            c1.metric(
                "Suvartojimas faile",
                f"{actual_load_total / 1000:.1f} MWh"
            )

            c2.metric(
                "Intervalas",
                f"{dt_hours * 60:.0f} min."
            )

            c3.metric(
                "Didžiausia apkrova",
                f"{peak_load_kw:.1f} kW"
            )

            c4.metric(
                "Įrašų skaičius",
                f"{len(load_df):,}"
            )

            st.write(
                "Laikotarpis:",
                first_timestamp,
                "–",
                last_timestamp
            )

            if price_coverage_pct is not None:
                st.metric(
                    "Faktinių kainų padengimas",
                    f"{price_coverage_pct:.1f} %"
                )

        # ----------------------------------------------------
        # SEZONINĖ ANALIZĖ IR 24 H AUTONOMIJA
        # ----------------------------------------------------
        result_sizing = core.calculate_sizing_metrics(
            base_hourly,
            pv_kw=summary["pv_kw"],
            discharge_eff=0.95,
            soc_min=0.05,
            soc_max=0.95,
            summer_reserve_pct=15.0,
            cfg=MODEL_CONFIG
        )

        st.markdown("### Sezoninis kaupimo sistemos vertinimas")
        r1, r2, r3, r4 = st.columns(4)
        r1.metric(
            "Vasaros dienos PV padengimas",
            f"{result_sizing['summer_day_pv_coverage_pct']:.1f} %"
        )
        r2.metric(
            "Dienos, kai PV pertekliaus pakanka nakčiai",
            f"{result_sizing['summer_full_charge_days_pct']:.1f} %"
        )
        r3.metric(
            "Žiemos PV padengimas",
            f"{result_sizing['winter_pv_coverage_pct']:.1f} %"
        )
        r4.metric(
            "Įkrovimas iš tinklo",
            f"{summary.get('grid_charge_kwh', 0.0) / 1000:.1f} MWh"
        )

        daylight = result_sizing["monthly_daylight"].copy()
        daylight["Mėnuo"] = daylight["month"].map({
            1: "Sau", 2: "Vas", 3: "Kov", 4: "Bal", 5: "Geg", 6: "Bir",
            7: "Lie", 8: "Rgp", 9: "Rgs", 10: "Spa", 11: "Lap", 12: "Gru"
        })
        daylight = daylight.set_index("Mėnuo")[["day_length_h", "night_length_h"]]
        daylight.columns = ["Dienos valandos", "Nakties valandos"]
        st.markdown("### Dienos ir nakties trukmė")
        st.line_chart(daylight)

        # ----------------------------------------------------
        # MĖNESINIAI REZULTATAI
        # ----------------------------------------------------

        monthly = (
            hourly_result
            .set_index("datetime")
            [
                [
                    "load_kwh",
                    "pv_kwh",
                    "grid_import_kwh",
                    "export_kwh",
                    "bess_charge_kwh",
                    "bess_discharge_kwh"
                ]
            ]
            .resample("ME")
            .sum()
        )

        st.markdown(
            "### Mėnesinis energijos balansas"
        )

        chart_data = monthly[
            [
                "load_kwh",
                "pv_kwh",
                "grid_import_kwh"
            ]
        ]

        st.line_chart(
            chart_data
        )

        # ----------------------------------------------------
        # VIENOS DIENOS ANALIZĖ
        # ----------------------------------------------------

        st.markdown(
            "### Vienos dienos analizė"
        )

        min_date = hourly_result[
            "datetime"
        ].min().date()

        max_date = hourly_result[
            "datetime"
        ].max().date()

        selected_date = st.date_input(
            "Pasirinkite datą",
            value=min_date,
            min_value=min_date,
            max_value=max_date
        )

        day_data = hourly_result[
            hourly_result[
                "datetime"
            ].dt.date == selected_date
        ]

        if len(day_data) > 0:

            day_chart = day_data.set_index(
                "datetime"
            )[
                [
                    "load_kwh",
                    "pv_kwh",
                    "soc_kwh"
                ]
            ]

            st.line_chart(
                day_chart
            )

        # ----------------------------------------------------
        # ATSISIUNTIMAS
        # ----------------------------------------------------

        csv = (
            hourly_result
            .to_csv(index=False)
            .encode("utf-8")
        )

        st.download_button(
            "Atsisiųsti valandinius rezultatus CSV",
            data=csv,
            file_name="PV_BESS_valandiniai_rezultatai.csv",
            mime="text/csv"
        )

        # Pilnas Excel eksportas: scenarijus + mėnesiai + sezonai + sizing +
        # valandiniai rezultatai + optimizacija (jei ji jau paleista).
        assumptions = {
            "PV profilio šaltinis": pv_profile_source_used,
            "Metinė specifinė PV generacija, kWh/kW": float(annual_specific_yield),
            "Referencinė PV galia, kW": float(reference_pv_kw),
            "Importo galios riba, kW": float(grid_import_limit_kw),
            "Eksporto galios riba, kW": float(grid_export_limit_kw),
            "Pakaitinė elektros kaina, €/kWh": float(price_fallback),
            "Faktinių kainų padengimas, %": price_coverage_pct,
            "PV galia, kW": summary.get("pv_kw"),
            "BESS talpa, kWh": summary.get("bess_kwh"),
            "BESS galia, kW": summary.get("bess_kw"),
            "Naudota degradacijos kaina, €/kWh": summary.get("degradation_cost_eur_kwh_used"),
            "Minimali arbitražo marža, €/kWh": summary.get("minimum_arbitrage_margin_eur_kwh_used"),
            "Vasaros mėnesiai": "4-10",
            "Pereinamieji mėnesiai": "3, 11",
            "Žiemos mėnesiai": "12, 1, 2",
        }

        validation = core.validate_energy_balance(
            hourly_result,
            charge_eff=0.95,
            discharge_eff=0.95
        )

        optimization_for_export = st.session_state.get("optimization_results")
        best_for_export = st.session_state.get("best_scenario")

        excel_bytes = build_excel_export(
            summary=summary,
            hourly_result=hourly_result,
            monthly=monthly,
            result_sizing=result_sizing,
            daylight=daylight,
            assumptions=assumptions,
            optimization_results=optimization_for_export,
            best_scenario=best_for_export,
            validation=validation,
        )

        st.download_button(
            "Atsisiųsti pilną modelio rezultatų Excel",
            data=excel_bytes,
            file_name="PV_BESS_modelio_rezultatai.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            type="primary"
        )

    else:

        st.info(
            "Dar nėra apskaičiuoto vieno scenarijaus."
        )

    # ========================================================
    # OPTIMIZAVIMO REZULTATAI
    # ========================================================

    if (
        "optimization_results"
        in st.session_state
    ):

        st.markdown(
            "---"
        )

        st.markdown(
            "## Optimizavimo rezultatai"
        )

        optimization_results = (
            st.session_state[
                "optimization_results"
            ]
        )

        if (
            "best_scenario"
            in st.session_state
            and
            st.session_state[
                "best_scenario"
            ]
            is not None
        ):

            best = st.session_state[
                "best_scenario"
            ]

            st.success(
                "Rastas reikalavimus atitinkantis optimalus variantas."
            )

            c1, c2, c3, c4 = st.columns(4)

            c1.metric(
                "Optimali PV galia",
                f"{best['pv_kw']:.0f} kW"
            )

            c2.metric(
                "Optimali BESS talpa",
                f"{best['bess_kwh']:.0f} kWh"
            )

            c3.metric(
                "Optimali BESS galia",
                f"{best['bess_kw']:.0f} kW"
            )

            c4.metric(
                "Savarankiškumas",
                f"{best['self_sufficiency_pct']:.1f} %"
            )

            c5, c6, c7, c8 = st.columns(4)

            c5.metric(
                "CAPEX",
                f"{best['total_capex_eur']:,.0f} €"
            )

            c6.metric(
                "NPC",
                f"{best['npc_eur']:,.0f} €"
            )

            c7.metric(
                "NPV",
                f"{best['npv_eur']:,.0f} €"
            )

            if pd.notna(
                best["simple_payback_years"]
            ):

                payback_text = (
                    f"{best['simple_payback_years']:.1f} m."
                )

            else:

                payback_text = "–"

            c8.metric(
                "Atsipirkimas",
                payback_text
            )

        else:

            st.warning(
                "Nė vienas scenarijus nepasiekė "
                "pasirinkto savarankiškumo reikalavimo."
            )

        # ====================================================
        # TOP 20
        # ====================================================

        st.markdown(
            "### 20 mažiausio NPC scenarijų"
        )

        top20 = (
            optimization_results
            .sort_values(
                "npc_eur"
            )
            .head(20)
            [
                [
                    "pv_kw",
                    "bess_kwh",
                    "bess_kw",
                    "bess_duration_h",
                    "self_sufficiency_pct",
                    "self_consumption_pct",
                    "grid_import_kwh",
                    "export_kwh",
                    "total_capex_eur",
                    "npc_eur",
                    "npv_eur",
                    "simple_payback_years"
                ]
            ]
            .round(2)
        )

        st.dataframe(
            top20,
            use_container_width=True
        )

        optimization_csv = (
            optimization_results
            .to_csv(
                index=False
            )
            .encode("utf-8")
        )

        st.download_button(
            "Atsisiųsti visus optimizavimo rezultatus CSV",
            data=optimization_csv,
            file_name="PV_BESS_optimizavimo_rezultatai.csv",
            mime="text/csv"
        )


# ============================================================
# 14. DUOMENŲ KOKYBĖS INFORMACIJA
# ============================================================

st.sidebar.markdown(
    "---"
)

st.sidebar.subheader(
    "Duomenų būsena"
)

if pv_profile_source_used:
    st.sidebar.success(
        f"PV profilis: {pv_profile_source_used}"
    )
else:
    st.sidebar.warning(
        "PV profilis: dar nesukurtas"
    )

if load_source == "Modeliuotas profilis":

    st.sidebar.warning(
        "Vartojimas: modeliuotas"
    )

elif load_file is not None:

    st.sidebar.success(
        "Vartojimas: faktinis"
    )

else:

    st.sidebar.warning(
        "Vartojimas: neįkeltas"
    )

if price_source == "Faktinės NPS kainos" and price_file is not None:
    st.sidebar.success("Elektros kaina: faktinės NPS kainos")
else:
    st.sidebar.warning("Elektros kaina: fiksuota / pakaitinė")

st.sidebar.info(
    "Ekonominės prielaidos šiuo metu "
    "įvedamos vartotojo ir turi būti "
    "patikrintos pagal konkretaus projekto duomenis."
)

